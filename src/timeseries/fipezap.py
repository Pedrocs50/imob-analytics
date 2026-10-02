from __future__ import annotations

import unicodedata
from datetime import date, datetime, timedelta

import openpyxl

from src.timeseries.config import DB_PATH, FIPEZAP_XLSX
from src.timeseries.storage import conectar, gravar_observacoes, registrar_coleta, registrar_serie

# abas lidas do Excel do FipeZAP -> (escopo, rotulo). Jacarei nao consta na planilha;
# Sao Jose dos Campos e a cidade coberta mais proxima.
ABAS = {
    "Índice FipeZAP": ("nacional", "Indice FipeZAP (composto)"),
    "São José dos Campos": ("sjc", "Sao Jose dos Campos"),
}

# Guardamos so o que nao e derivavel. Var. mensal e var. em 12 meses saem do
# numero-indice e podem ser recalculadas quando necessario.
METRICAS = {
    "numero-indice": ("indice", "indice", "Numero-indice"),
    "preco medio (r$/m2)": ("preco_m2", "R$/m2", "Preco medio"),
    "(% - mensalizada)": ("yield", "% a.m.", "Rentabilidade do aluguel (yield mensalizado)"),
}

# O indice do mes t e calculado com anuncios do proprio mes e publicado no mes seguinte
# (evidencias em docs/SERIES_TEMPORAIS.md); na convencao dos paineis a defasagem e 1.
DEFASAGEM_FIPEZAP = 1

LINHA_PRIMEIRO_DADO = 5
COL_DATA = 2  # coluna B


def _norm(texto) -> str:
    if texto is None:
        return ""
    sem_acento = unicodedata.normalize("NFKD", str(texto)).encode("ascii", "ignore").decode("ascii")
    return " ".join(sem_acento.lower().split())


def _para_data(valor) -> date | None:
    if isinstance(valor, datetime):
        return valor.date().replace(day=1)
    if isinstance(valor, date):
        return valor.replace(day=1)
    if isinstance(valor, (int, float)):  # data serial do Excel
        return (date(1899, 12, 30) + timedelta(days=int(valor))).replace(day=1)
    return None


def _preencher_a_direita(linha: tuple) -> list:
    """Celulas mescladas chegam vazias; propaga o ultimo rotulo para a direita."""
    saida, atual = [], None
    for celula in linha:
        if celula not in (None, ""):
            atual = celula
        saida.append(atual)
    return saida


def _cabecalhos(linhas: list[tuple]) -> dict[int, dict]:
    """Descreve cada coluna de dados a partir das 4 linhas de cabecalho."""
    segmento = _preencher_a_direita(linhas[0])
    operacao = _preencher_a_direita(linhas[1])
    metrica = _preencher_a_direita(linhas[2])
    dorm = linhas[3]

    colunas: dict[int, dict] = {}
    for i in range(COL_DATA, len(dorm)):
        if dorm[i] in (None, ""):
            continue
        colunas[i] = {
            "segmento": _norm(segmento[i]),
            "operacao": _norm(operacao[i]),
            "metrica": _norm(metrica[i]),
            "dorm": str(dorm[i]).strip(),
        }
    return colunas


def _id_serie(escopo: str, col: dict) -> tuple[str, dict] | None:
    metrica = METRICAS.get(col["metrica"])
    if metrica is None:
        return None
    codigo_metrica, unidade, rotulo_metrica = metrica

    segmento = "res" if "residenciais" in col["segmento"] else "com" if "comerciais" in col["segmento"] else None
    if segmento is None:
        return None

    if codigo_metrica == "yield":
        operacao, nome_operacao = "aluguel", "aluguel"
    elif col["operacao"].startswith("venda"):
        operacao, nome_operacao = "venda", "venda"
    elif col["operacao"].startswith("locacao"):
        operacao, nome_operacao = "locacao", "locacao"
    else:
        return None

    sufixo = "" if col["dorm"].lower() == "total" else f"_{col['dorm'].lower()}"
    partes = [f"fipezap_{escopo}_{segmento}", operacao, codigo_metrica]
    serie_id = "_".join(partes) + sufixo
    detalhe = "total" if not sufixo else f"{col['dorm']} dormitorio(s)"
    nome = f"FipeZAP {rotulo_metrica} - {'residencial' if segmento == 'res' else 'comercial'} {nome_operacao} ({detalhe})"
    return serie_id, {"unidade": unidade, "tipo": "indice" if codigo_metrica == "indice" else "nivel", "nome": nome}


def ler_fipezap(caminho: str = FIPEZAP_XLSX) -> dict[str, dict]:
    """Le o Excel do FipeZAP. Retorna {serie_id: {meta..., "pontos": [(data, valor)]}}."""
    livro = openpyxl.load_workbook(caminho, read_only=True, data_only=True)
    resultado: dict[str, dict] = {}
    try:
        for aba, (escopo, rotulo_escopo) in ABAS.items():
            if aba not in livro.sheetnames:
                raise KeyError(f"aba {aba!r} nao encontrada em {caminho}")
            linhas = list(livro[aba].iter_rows(values_only=True))
            colunas = _cabecalhos(linhas[:4])

            for i, col in colunas.items():
                descricao = _id_serie(escopo, col)
                if descricao is None:
                    continue
                serie_id, meta = descricao
                # o Excel guarda o yield como fracao (0,004); padronizamos em % a.m.
                fator = 100.0 if meta["unidade"] == "% a.m." else 1.0
                pontos = []
                for linha in linhas[LINHA_PRIMEIRO_DADO - 1:]:
                    if len(linha) <= i:
                        continue
                    data = _para_data(linha[COL_DATA - 1])
                    valor = linha[i]
                    # '.' marca periodo sem dado; textos de rodape nao tem data valida
                    if data is None or not isinstance(valor, (int, float)):
                        continue
                    pontos.append((data, float(valor) * fator))
                if pontos:
                    resultado[serie_id] = {
                        **meta,
                        "nome": f"{meta['nome']} - {rotulo_escopo}",
                        "escopo": escopo,
                        "pontos": sorted(pontos),
                    }
    finally:
        livro.close()
    return resultado


def importar_fipezap(caminho: str = FIPEZAP_XLSX, db_path: str = DB_PATH) -> dict[str, int]:
    """Importa o Excel do FipeZAP para o banco de series temporais."""
    series = ler_fipezap(caminho)
    contagem: dict[str, int] = {}
    with conectar(db_path) as conexao:
        for serie_id, dados in series.items():
            registrar_serie(
                conexao,
                {
                    "id": serie_id, "fonte": "FIPEZAP", "codigo": None, "nome": dados["nome"],
                    "unidade": dados["unidade"], "tipo": dados["tipo"], "frequencia": "M",
                    "defasagem_meses": DEFASAGEM_FIPEZAP, "escopo": dados["escopo"],
                },
            )
            n = gravar_observacoes(conexao, serie_id, dados["pontos"])
            registrar_coleta(conexao, serie_id, n, "ok", arquivo_bruto=caminho)
            contagem[serie_id] = n
    print(f"[SERIES] FipeZAP: {len(contagem)} series importadas de {caminho}")
    return contagem
