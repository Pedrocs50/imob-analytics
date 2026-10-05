from __future__ import annotations

import contextlib
import io
import os
import time
import warnings

import matplotlib

matplotlib.use("Agg")
import matplotlib.pyplot as plt
import numpy as np
import openpyxl
import pandas as pd
from statsmodels.tools.sm_exceptions import ConvergenceWarning
from statsmodels.tsa.arima.model import ARIMA
from statsmodels.tsa.stattools import grangercausalitytests

from src.analysis.series_eda import preditores_estacionarios
from src.analysis.utils import md, salvar_fig
from src.timeseries.arima import diebold_mariano
from src.timeseries.config import FIPEZAP_XLSX
from src.timeseries.fipezap import COL_DATA, LINHA_PRIMEIRO_DADO, _cabecalhos, _para_data
from src.timeseries.panel import carregar_wide, montar_paineis

REPORT = os.path.join("reports", "results", "arima_plus.md")
TAB_DIR = os.path.join("reports", "results", "arima")
HORIZONTES = (1, 3, 6, 12)
LAG_EXOG = 12  # preditores defasados 12 meses: conhecidos ate o horizonte de 12 meses
ALVOS = {
    "nacional (nominal)": ("fipezap_nacional_res_venda_preco_m2", 84),
    "SJC (nominal)": ("fipezap_sjc_res_venda_preco_m2", 60),
}
ORDEM = (0, 2, 3)  # escolhida na etapa anterior (docs/MODELOS_TEMPORAIS.md)
CONJUNTOS_EXOG = {
    "ARIMAX credito": ["bcb_credimob_saldo"],
    "ARIMAX juros e confianca": ["bcb_selic_meta", "ipea_icc"],
    "ARIMAX credito + juros + confianca + IBC-Br": ["bcb_credimob_saldo", "bcb_selic_meta", "ipea_icc", "bcb_ibcbr"],
}
REFERENCIA = "ARIMA(0,2,3)"
TEND = "Tendencia (12 meses)"
COMBO = "Combinacao ARIMA + tendencia"


# ----------------------------------------------------------------------------- previsao
def _arima_log(y: pd.Series) -> tuple[np.ndarray, np.ndarray, np.ndarray]:
    with warnings.catch_warnings():
        warnings.simplefilter("ignore")
        res = ARIMA(y, order=ORDEM, trend="n").fit()
        fc = res.get_forecast(steps=max(HORIZONTES))
        ic = fc.conf_int(alpha=0.05)
    return fc.predicted_mean.to_numpy(), ic.iloc[:, 0].to_numpy(), ic.iloc[:, 1].to_numpy()


def _arimax_log(y: pd.Series, x: pd.DataFrame) -> np.ndarray:
    """ARMA(0,3) em w = d2 log(y) com preditores defasados, e reintegracao para o nivel.
    x ja esta defasado em LAG_EXOG meses e indexado pela data da observacao a prever."""
    w = y.diff().diff().dropna()
    xtr = x.reindex(w.index)
    valido = xtr.notna().all(axis=1)
    w, xtr = w[valido], xtr[valido]
    media, desvio = xtr.mean(), xtr.std().replace(0, 1)
    futuro = x.reindex(pd.date_range(y.index[-1] + pd.offsets.MonthBegin(1), periods=max(HORIZONTES), freq="MS"))
    with warnings.catch_warnings():
        warnings.simplefilter("ignore", ConvergenceWarning)
        warnings.simplefilter("ignore")
        res = ARIMA(w.to_numpy(), exog=((xtr - media) / desvio).to_numpy(), order=(0, 0, 3), trend="n").fit()
        w_hat = res.forecast(steps=max(HORIZONTES), exog=((futuro - media) / desvio).to_numpy())
    cresc = (y.iloc[-1] - y.iloc[-2]) + np.cumsum(w_hat)  # d y_{T+h}
    return y.iloc[-1] + np.cumsum(cresc)


def walk_forward(y: pd.Series, n0: int, x_exog: dict[str, pd.DataFrame]) -> pd.DataFrame:
    n = len(y)
    linhas = []
    for i in range(n0 - 1, n - 1):
        passado = y.iloc[: i + 1]
        previsoes = {}
        arima, inf, sup = _arima_log(passado)
        previsoes[REFERENCIA] = arima
        c12 = passado.diff().iloc[-12:].mean()
        previsoes[TEND] = passado.iloc[-1] + np.arange(1, 13) * c12
        previsoes[COMBO] = (arima + previsoes[TEND]) / 2
        for nome, x in x_exog.items():
            try:
                previsoes[nome] = _arimax_log(passado, x)
            except Exception:
                previsoes[nome] = np.full(12, np.nan)
        for h in HORIZONTES:
            if i + h >= n:
                continue
            linha = {"origem": y.index[i], "h": h, "real": y.iloc[i + h], "ultimo": passado.iloc[-1],
                     "inf95": inf[h - 1], "sup95": sup[h - 1]}
            linha.update({m: v[h - 1] for m, v in previsoes.items()})
            linhas.append(linha)
    return pd.DataFrame(linhas)


# ----------------------------------------------------------------------------- avaliacao
def metricas(wf: pd.DataFrame, modelos: list[str]) -> pd.DataFrame:
    linhas = []
    ref = {}
    for h, g in wf.groupby("h"):
        real = np.exp(g["real"])
        for m in modelos:
            prev = np.exp(g[m])
            ok = prev.notna()
            erro = (real - prev)[ok]
            linhas.append({"h (meses)": h, "modelo": m, "n origens": int(ok.sum()), "MAPE (%)": float((erro.abs() / real[ok]).mean() * 100),
                           "RMSE": float(np.sqrt((erro ** 2).mean()))})
    t = pd.DataFrame(linhas)
    base = t[t["modelo"] == REFERENCIA].set_index("h (meses)")["RMSE"]
    t["RMSE relativo ao ARIMA"] = t["RMSE"] / t["h (meses)"].map(base)
    return t


def testes_dm(wf: pd.DataFrame, modelos: list[str]) -> pd.DataFrame:
    linhas = []
    for h in HORIZONTES:
        for m in modelos:
            if m == REFERENCIA:
                continue
            g = wf[(wf["h"] == h) & wf[m].notna()]
            gg = g.assign(arima=g[m])  # reaproveita a funcao: "arima" = modelo avaliado
            gg[REFERENCIA] = g[REFERENCIA]
            r = diebold_mariano(gg, h, REFERENCIA)
            linhas.append({"modelo": m, "contra": REFERENCIA, "h (meses)": h, "estatistica DM": r["estatistica DM"],
                           "p (bilateral)": r["p (bilateral)"], "n": r["n"]})
    return pd.DataFrame(linhas)


def cobertura(wf: pd.DataFrame) -> pd.DataFrame:
    linhas = []
    for h, g in wf.groupby("h"):
        dentro = (g["real"] >= g["inf95"]) & (g["real"] <= g["sup95"])
        largura = ((np.exp(g["sup95"]) - np.exp(g["inf95"])) / np.exp(g["real"]) * 100).mean()
        linhas.append({"h (meses)": h, "cobertura do IC 95% (%)": float(dentro.mean() * 100), "largura media do IC (% do preco)": float(largura), "n": len(g)})
    return pd.DataFrame(linhas)


def por_periodo(wf: pd.DataFrame, modelos: list[str]) -> pd.DataFrame:
    linhas = []
    for nome, filtro in (("origens ate 2019", wf["origem"] < "2020-01-01"), ("origens de 2020 em diante", wf["origem"] >= "2020-01-01")):
        sub = wf[filtro]
        for h, g in sub.groupby("h"):
            real = np.exp(g["real"])
            for m in (REFERENCIA, TEND, COMBO):
                linhas.append({"periodo": nome, "h (meses)": h, "modelo": m, "n": len(g),
                               "MAPE (%)": float(((real - np.exp(g[m])).abs() / real).mean() * 100)})
    t = pd.DataFrame(linhas)
    return t[t["h (meses)"].isin([3, 12])]


# ----------------------------------------------------------------------------- cidades vizinhas
def ler_cidades(caminho: str = FIPEZAP_XLSX) -> pd.DataFrame:
    """Preco medio de venda residencial (total) de cada cidade do Excel do FipeZAP."""
    wb = openpyxl.load_workbook(caminho, read_only=True, data_only=True)
    series = {}
    try:
        for nome in wb.sheetnames:
            if nome in ("Resumo", "Aux", "Índice FipeZAP"):
                continue
            linhas = list(wb[nome].iter_rows(values_only=True))
            try:
                cols = _cabecalhos(linhas[:4])
            except Exception:
                continue
            alvo = [i for i, c in cols.items() if c["segmento"].startswith("imoveis residenciais")
                    and c["operacao"].startswith("venda") and c["metrica"].startswith("preco") and c["dorm"].lower() == "total"]
            if not alvo:
                continue
            i = alvo[0]
            pts = {}
            for r in linhas[LINHA_PRIMEIRO_DADO - 1:]:
                d = _para_data(r[COL_DATA - 1]) if len(r) >= COL_DATA else None
                if d is not None and len(r) > i and isinstance(r[i], (int, float)):
                    pts[pd.Timestamp(d)] = float(r[i])
            if len(pts) >= 100:
                series[nome] = pd.Series(pts).sort_index()
    finally:
        wb.close()
    return pd.DataFrame(series)


def cidades_defasadas(precos: pd.DataFrame, alvo: pd.Series, max_lag: int = 6) -> pd.DataFrame:
    """Alguma cidade antecipa o alvo? Correlacao do alvo (d2 do log) com o d2 do log de cada cidade
    defasado de 1 a 6 meses, e teste de Granger. Exploratorio, com varios testes: usar com cautela."""
    def d2(s: pd.Series) -> pd.Series:
        return np.log(s).diff().diff()

    y = d2(alvo)
    linhas = []
    for cidade in precos.columns:
        x = d2(precos[cidade])
        par = pd.concat([y, x], axis=1).dropna()
        if len(par) < 60:
            continue
        corrs = {k: y.corr(x.shift(k)) for k in range(1, max_lag + 1)}
        with warnings.catch_warnings(), contextlib.redirect_stdout(io.StringIO()):
            warnings.simplefilter("ignore")
            res = grangercausalitytests(par.to_numpy(), maxlag=max_lag)
        p = {k: float(res[k][0]["ssr_ftest"][1]) for k in range(1, max_lag + 1)}
        k_p = min(p, key=p.get)
        k_r = max(corrs, key=lambda k: abs(corrs[k]))
        linhas.append({"cidade": cidade, "n": len(par), "maior |r| (lag 1-6)": abs(corrs[k_r]), "lag do maior |r|": k_r,
                       "menor p (Granger)": p[k_p], "lag do menor p": k_p})
    t = pd.DataFrame(linhas).sort_values("menor p (Granger)").reset_index(drop=True)
    t.attrs["bonferroni"] = 0.05 / (len(t) * max_lag) if len(t) else np.nan
    return t


# ----------------------------------------------------------------------------- figura e relatorio
def figura(tabelas: dict[str, pd.DataFrame], modelos: list[str]) -> str:
    fig, eixos = plt.subplots(1, len(tabelas), figsize=(7 * len(tabelas), 4.6), sharey=True)
    for eixo, (nome, t) in zip(np.atleast_1d(eixos), tabelas.items()):
        piv = t.pivot(index="h (meses)", columns="modelo", values="RMSE relativo ao ARIMA")[[m for m in modelos if m != REFERENCIA]]
        piv.plot.bar(ax=eixo, rot=0, width=0.85)
        eixo.axhline(1, color="k", ls="--", lw=1)
        eixo.set_title(f"{nome}: RMSE relativo ao ARIMA(0,2,3)", fontsize=10)
        eixo.set_xlabel("horizonte (meses)")
        eixo.grid(axis="y", alpha=0.3)
        eixo.legend(fontsize=7)
    np.atleast_1d(eixos)[0].set_ylabel("RMSE / RMSE do ARIMA (menor e melhor)")
    fig.tight_layout()
    return salvar_fig(fig, "arima_plus_comparacao.png")


def executar() -> dict:
    inicio = time.time()
    wide, meta = carregar_wide()
    p = montar_paineis(wide, meta)
    ref = p.mensal_referencia
    x_est, _ = preditores_estacionarios(ref, p.defasagem_mensal)
    modelos = [REFERENCIA, TEND, COMBO, *CONJUNTOS_EXOG]
    os.makedirs(TAB_DIR, exist_ok=True)

    res = {}
    for nome, (col, n0) in ALVOS.items():
        y = np.log(ref[col].dropna())
        y.index = pd.DatetimeIndex(y.index, freq="MS")
        x_exog = {k: x_est[cols].shift(LAG_EXOG) for k, cols in CONJUNTOS_EXOG.items()}
        t0 = time.time()
        wf = walk_forward(y, n0, x_exog)
        res[nome] = {"met": metricas(wf, modelos), "dm": testes_dm(wf, modelos), "cob": cobertura(wf), "per": por_periodo(wf, modelos)}
        slug = nome.replace(" ", "_").replace("(", "").replace(")", "")
        for k, df in res[nome].items():
            df.to_csv(os.path.join(TAB_DIR, f"arima_plus_{slug}_{k}.csv"), index=False, encoding="utf-8")
        print(f"[ARIMA+] {nome}: {wf['origem'].nunique()} origens em {time.time() - t0:.0f}s")

    precos = ler_cidades()
    vizinhas = {}
    for nome, (col, _) in ALVOS.items():
        alvo = ref[col].dropna()
        outras = precos.drop(columns=[c for c in precos.columns if c.startswith("São José dos Campos")], errors="ignore")
        vizinhas[nome] = cidades_defasadas(outras, alvo)
        vizinhas[nome].to_csv(os.path.join(TAB_DIR, f"arima_plus_vizinhas_{nome.split()[0]}.csv"), index=False, encoding="utf-8")

    fig = figura({k: v["met"] for k, v in res.items()}, modelos)
    escrever_relatorio(res, vizinhas, [fig], time.time() - inicio, precos.shape[1])
    print(f"[ARIMA+] concluido em {time.time() - inicio:.0f}s")
    return {"res": res, "vizinhas": vizinhas}


def escrever_relatorio(res: dict, vizinhas: dict, figuras: list[str], segundos: float, n_cidades: int) -> None:
    L = ["# ARIMA: combinacao, ARIMAX, cidades vizinhas e avaliacao ampliada", "",
         "Gerado por `python main.py arima-plus`. Nao editar manualmente. Leitura em `docs/MODELOS_TEMPORAIS.md`.", "",
         f"Mesmo walk-forward do ARIMA (`docs/MODELOS_TEMPORAIS.md`), ordem {ORDEM} fixa, alvos nominais. Comparacao **contra o ARIMA(0,2,3)** "
         "(RMSE relativo menor que 1 = melhor que o ARIMA). Modelos:", "",
         f"- **{TEND}:** crescimento medio dos ultimos 12 meses; **{COMBO}:** media simples dos dois.",
         f"- **ARIMAX:** ARMA(0,3) em d2 do log do preco com preditores estacionarios **defasados {LAG_EXOG} meses** (conhecidos em todos os "
         "horizontes ate 12) e previsao reintegrada ao nivel. Os preditores usam a diferenciacao decidida na analise exploratoria.", "",
         f"Tempo total: {segundos:.0f} s.", ""]
    for nome, r in res.items():
        L += [f"## {nome}", "", "### Erro por horizonte (MAPE em %, RMSE em R$/m2 e RMSE relativo ao ARIMA)", "", md(r["met"].round(3), 3), "",
              "### Diebold-Mariano contra o ARIMA (estatistica negativa: o modelo e melhor que o ARIMA)", "", md(r["dm"].round(3), 3), "",
              "### Cobertura e largura dos intervalos de previsao do ARIMA (nominal 95%)", "", md(r["cob"].round(2), 2), "",
              "### Erro por periodo (MAPE em %; 3 e 12 meses)", "", md(r["per"].round(3), 3), ""]
    for nome, t in vizinhas.items():
        L += [f"## Cidades que antecipam o FipeZAP de {nome}? ({n_cidades} cidades; d2 do log, lags de 1 a 6)", "",
              f"Limiar de Bonferroni: p < {t.attrs.get('bonferroni', float('nan')):.5f}. Dez menores p:", "", md(t.head(10).round(4), 4), ""]
    L += ["## Figuras", ""] + [f"- `{f}`" for f in figuras]
    os.makedirs(os.path.dirname(REPORT), exist_ok=True)
    with open(REPORT, "w", encoding="utf-8") as arquivo:
        arquivo.write("\n".join(L) + "\n")
